import argparse
import json
from pathlib import Path

import openpyxl


def convert_xlsx_to_layout(xlsx_path: str | Path, output_path: str | Path, store_id: str = "ST1008") -> Path:
    xlsx_path = Path(xlsx_path)
    output_path = Path(output_path)
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    zones = []
    for row in sheet.iter_rows(values_only=True):
        values = [value for value in row if value is not None]
        if len(values) >= 5 and str(values[0]).lower() in {"zone", "billing", "entrance"}:
            zone_id = str(values[1]).strip().lower().replace(" ", "_")
            points = _parse_points(str(values[-1]))
            if points:
                zones.append({"zone_id": zone_id, "name": str(values[1]), "type": str(values[0]), "polygon": points})

    if not zones:
        zones = [
            {"zone_id": "entrance", "name": "Entrance", "type": "entry", "polygon": [[40, 520], [1240, 520], [1240, 720], [40, 720]]},
            {"zone_id": "aisle_skin", "name": "Skin Care Aisle", "type": "shopping", "polygon": [[80, 160], [520, 160], [520, 500], [80, 500]]},
            {"zone_id": "aisle_makeup", "name": "Makeup Aisle", "type": "shopping", "polygon": [[560, 160], [1040, 160], [1040, 500], [560, 500]]},
            {"zone_id": "billing", "name": "Billing Counter", "type": "billing", "polygon": [[900, 500], [1240, 500], [1240, 720], [900, 720]]},
        ]

    layout = {
        "store_id": store_id,
        "camera_defaults": {"entry_line": [[80, 620], [1180, 620]], "exit_line": [[80, 680], [1180, 680]]},
        "queue_threshold": 3,
        "zones": zones,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(layout, indent=2), encoding="utf-8")
    return output_path


def _parse_points(value: str) -> list[list[int]]:
    points = []
    for token in value.replace(";", " ").split():
        if "," not in token:
            continue
        x, y = token.split(",", 1)
        if x.strip().isdigit() and y.strip().isdigit():
            points.append([int(x), int(y)])
    return points


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx_path")
    parser.add_argument("--output", default="data/store_layout.json")
    parser.add_argument("--store-id", default="ST1008")
    args = parser.parse_args()
    print(convert_xlsx_to_layout(args.xlsx_path, args.output, args.store_id))


if __name__ == "__main__":
    main()
